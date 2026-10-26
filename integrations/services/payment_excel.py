"""生成付款申请Excel报表。

使用 openpyxl 生成包含4个工作表的Excel报表:
1. 汇总统计
2. 按收款人汇总
3. 详细清单
4. 原始JSON数据
"""

import json
import logging
from decimal import Decimal

logger = logging.getLogger("integrations")


def _format_money(amount):
    """格式化金额为千分位格式。"""
    if amount is None:
        return '¥0.00'
    return f"¥{float(amount):,.2f}"


def _format_status(status):
    """将状态码转换为中文显示。"""
    status_map = {
        'COMPLETED': '已完成',
        'RUNNING': '进行中',
        'TERMINATED': '已终止',
        'NEW': '新建',
    }
    return status_map.get(status, status)


def generate_payment_excel(data, period_info, filename):
    """生成付款申请Excel报表。

    Args:
        data: analyze_payment_applications 返回的数据
        period_info: 周期信息 {label, start_date, end_date}
        filename: 输出文件路径

    Returns:
        str: 生成的文件路径
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError as exc:
        raise RuntimeError("未安装 openpyxl,无法生成Excel。请运行: pip install openpyxl") from exc

    workbook = Workbook()

    # 删除默认sheet
    if 'Sheet' in workbook.sheetnames:
        del workbook['Sheet']

    # ========== 工作表1:汇总统计 ==========
    summary_sheet = workbook.create_sheet('汇总统计')
    summary_sheet.column_dimensions['A'].width = 25
    summary_sheet.column_dimensions['B'].width = 35

    # 表头样式
    header_fill = PatternFill(start_color='E0E0E0', end_color='E0E0E0', fill_type='solid')
    header_font = Font(bold=True)
    center_align = Alignment(horizontal='center', vertical='center')

    summary_sheet['A1'] = '统计项'
    summary_sheet['B1'] = '数值'
    summary_sheet['A1'].fill = header_fill
    summary_sheet['B1'].fill = header_fill
    summary_sheet['A1'].font = header_font
    summary_sheet['B1'].font = header_font
    summary_sheet['A1'].alignment = center_align
    summary_sheet['B1'].alignment = center_align

    # 数据行
    summary_rows = [
        ['统计周期', period_info['label']],
        ['起始日期', period_info['start_date'].strftime('%Y-%m-%d')],
        ['结束日期', period_info['end_date'].strftime('%Y-%m-%d')],
        ['总申请数', data['summary']['total_count']],
        ['总金额', _format_money(data['summary']['total_amount'])],
        ['已完成', data['summary']['by_status']['COMPLETED']],
        ['进行中', data['summary']['by_status']['RUNNING']],
        ['已终止', data['summary']['by_status']['TERMINATED']],
        ['新建', data['summary']['by_status']['NEW']],
    ]

    for idx, row in enumerate(summary_rows, start=2):
        summary_sheet[f'A{idx}'] = row[0]
        summary_sheet[f'B{idx}'] = row[1]

    # ========== 工作表2:按收款人汇总 ==========
    recipient_sheet = workbook.create_sheet('按收款人汇总')
    recipient_sheet.column_dimensions['A'].width = 40
    recipient_sheet.column_dimensions['B'].width = 12
    recipient_sheet.column_dimensions['C'].width = 20
    recipient_sheet.column_dimensions['D'].width = 12
    recipient_sheet.column_dimensions['E'].width = 12
    recipient_sheet.column_dimensions['F'].width = 12

    # 表头
    headers = ['收款人', '申请数', '总金额', '已完成', '进行中', '已终止']
    for col_idx, header in enumerate(headers, start=1):
        cell = recipient_sheet.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align

    # 数据行
    for row_idx, recipient_data in enumerate(data['by_recipient'], start=2):
        completed = sum(1 for d in recipient_data['details'] if d['status'] == 'COMPLETED')
        running = sum(1 for d in recipient_data['details'] if d['status'] == 'RUNNING')
        terminated = sum(1 for d in recipient_data['details'] if d['status'] == 'TERMINATED')

        recipient_sheet.cell(row=row_idx, column=1, value=recipient_data['recipient'])
        recipient_sheet.cell(row=row_idx, column=2, value=recipient_data['count'])
        recipient_sheet.cell(row=row_idx, column=3, value=_format_money(recipient_data['amount']))
        recipient_sheet.cell(row=row_idx, column=4, value=completed)
        recipient_sheet.cell(row=row_idx, column=5, value=running)
        recipient_sheet.cell(row=row_idx, column=6, value=terminated)

    # ========== 工作表3:详细清单 ==========
    detail_sheet = workbook.create_sheet('详细清单')
    detail_sheet.column_dimensions['A'].width = 15
    detail_sheet.column_dimensions['B'].width = 40
    detail_sheet.column_dimensions['C'].width = 18
    detail_sheet.column_dimensions['D'].width = 50
    detail_sheet.column_dimensions['E'].width = 18
    detail_sheet.column_dimensions['F'].width = 12
    detail_sheet.column_dimensions['G'].width = 15
    detail_sheet.column_dimensions['H'].width = 35
    detail_sheet.column_dimensions['I'].width = 15

    # 表头
    detail_headers = [
        '申请日期', '收款人', '付款金额', '付款事由',
        '期望付款日期', '状态', '申请人', '审批标题', '审批实例ID'
    ]
    for col_idx, header in enumerate(detail_headers, start=1):
        cell = detail_sheet.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align

    # 数据行
    row_idx = 2
    for recipient_data in data['by_recipient']:
        for detail in recipient_data['details']:
            detail_sheet.cell(row=row_idx, column=1, value=detail['create_date'])
            detail_sheet.cell(row=row_idx, column=2, value=recipient_data['recipient'])
            detail_sheet.cell(row=row_idx, column=3, value=_format_money(detail['amount']))
            detail_sheet.cell(row=row_idx, column=4, value=detail['reason'])
            detail_sheet.cell(row=row_idx, column=5, value=detail['expected_date'])
            detail_sheet.cell(row=row_idx, column=6, value=_format_status(detail['status']))
            detail_sheet.cell(row=row_idx, column=7, value=detail['originator'])
            detail_sheet.cell(row=row_idx, column=8, value=detail['title'])
            detail_sheet.cell(row=row_idx, column=9, value=detail['process_instance_id'])
            row_idx += 1

    # ========== 工作表4:原始JSON数据 ==========
    raw_sheet = workbook.create_sheet('原始数据')
    raw_sheet.column_dimensions['A'].width = 120

    # 将数据转为JSON字符串(处理Decimal和datetime类型)
    def json_default(obj):
        if isinstance(obj, Decimal):
            return float(obj)
        if hasattr(obj, 'isoformat'):  # datetime, date, time
            return obj.isoformat()
        raise TypeError(f"Object of type {type(obj).__name__} is not JSON serializable")

    json_str = json.dumps(
        {
            'period': period_info,
            'data': data,
        },
        ensure_ascii=False,
        indent=2,
        default=json_default,
    )
    raw_sheet['A1'] = json_str

    # 保存文件
    workbook.save(filename)
    logger.info(f"Excel报表已生成: {filename}")

    return filename

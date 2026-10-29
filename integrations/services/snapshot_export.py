"""统一采销日报 Excel：金蝶汇总 / 逐日金额，另附原有手工月度对比表。"""

from io import BytesIO
from decimal import Decimal
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from integrations.services.operations_summary import (
    snapshot_operations,
    snapshot_detail_context,
)

COMPARISON_SHEET = "采购入库金额与销售金额对比表"


def _number(value):
    number = Decimal(str(value or 0))
    if not number.is_finite():
        raise ValueError("日报金额必须为有限数值。")
    return float(number)


def _style(ws, header_row):
    ws.freeze_panes = f"B{header_row+1}"
    ws.column_dimensions["A"].width = 24
    for key in ("B", "C", "D", "E"):
        ws.column_dimensions[key].width = 24
    for cell in ws[header_row]:
        cell.fill = PatternFill("solid", fgColor="243C5A")
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center")
    for row in ws.iter_rows(min_row=header_row + 1):
        for cell in row:
            if isinstance(cell.value, (float, int)):
                cell.number_format = "#,##0.00;[Red]-#,##0.00"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = f"1:{header_row}"


def _daily_sheet(wb, title, rows, names):
    ws = wb.create_sheet(title)
    ws.append(
        ["日期"]
        + [("采购入库金额" if name == "purchase" else "销售金额") for name in names]
    )
    for row in rows:
        ws.append(
            [row["date"]] + [_number(row.get(name + "_amount")) for name in names]
        )
    _style(ws, 1)
    return ws


def _detail_sheet(wb, snapshot, title, source, headers, records, empty_message, total):
    ws = wb.create_sheet(title)
    ws.append([title])
    ws.append(["报表日期", snapshot.report_date.isoformat()])
    ws.append(["来源", source])
    ws.append(headers)
    for row in records:
        ws.append(row)
    if not records:
        ws.append([empty_message])
    ws.append(["合计"] + [None] * (len(headers) - 4) + total)
    _style(ws, 4)
    ws.column_dimensions["C"].width = 30
    ws.column_dimensions["D"].width = 30
    ws.column_dimensions["E"].width = 20
    ws.column_dimensions["F"].width = 24
    ws.column_dimensions["G"].width = 24
    for row in ws.iter_rows(min_row=5):
        row[4].number_format = "#,##0.######;[Red]-#,##0.######"
    ws["A3"].alignment = Alignment(wrap_text=False)
    ws["B3"].alignment = Alignment(wrap_text=True)
    ws.row_dimensions[3].height = 42
    ws.auto_filter.ref = f"A4:{ws.cell(ws.max_row - 1, len(headers)).coordinate}"
    return ws


def _create_detail_sheets(wb, snapshot, details):
    _detail_sheet(
        wb,
        snapshot,
        "日报提交明细",
        "人员填报日报 / 历史导入；仅报表当日，不以金蝶单据代填",
        ["业务类型", "日期", "客户/供应商", "提报人", "数量", "单位", "日报金额（元）"],
        [
            [
                row["business_label"],
                row.get("date", ""),
                row.get("party_name", ""),
                row.get("reporter", ""),
                _number(row["quantity"]) if row.get("quantity") is not None else None,
                row.get("unit") or "未记录单位",
                _number(row.get("amount")),
            ]
            for row in details["manual_rows"]
        ],
        "暂无提交记录（保留日报字段，未用金蝶数据代填）",
        [
            (
                _number(details["manual_quantity_total"])
                if details["manual_quantity_total"] is not None
                else None
            ),
            details["manual_quantity_unit"],
            _number(details["manual_total"]),
        ],
    )
    missing = details["kingdee_details_missing"]
    kingdee_sheet = _detail_sheet(
        wb,
        snapshot,
        "当日金蝶单据明细",
        "金蝶 K3Cloud 有效"
        + " / ".join(row["label"] for row in details["manual_ledgers"])
        + "单据；报表当日全量单据，不截断，退货保留负数。数量为源单据算术合计，混合单位未折算，不代表统一实物数量。",
        [
            "业务类型",
            "业务日期",
            "单据编号",
            "客户/供应商",
            "数量",
            "单位",
            "金蝶金额（元）",
        ],
        [
            [
                row["business_label"],
                row["date"],
                row["bill_no"],
                row.get("party_name", ""),
                _number(row["quantity"]) if row.get("quantity") is not None else None,
                row.get("unit") or "未记录单位",
                _number(row["amount"]),
            ]
            for row in details["kingdee_documents"]
        ],
        "历史快照未冻结金蝶单据明细，请重新生成" if missing else "当日暂无金蝶单据",
        [
            (
                _number(details["kingdee_quantity_total"])
                if not missing and details["kingdee_quantity_total"] is not None
                else None
            ),
            details["kingdee_quantity_unit"],
            None if missing else _number(details["kingdee_document_total"]),
        ],
    )

    kingdee_sheet.column_dimensions["C"].width = 38
    for row in kingdee_sheet.iter_rows(min_row=4, min_col=3, max_col=3):
        row[0].alignment = Alignment(wrap_text=False)


def export_snapshot_to_excel(snapshot) -> bytes:
    """只读取快照；导出/投递重试不会重新查询实时业务数据。"""
    data = snapshot_operations(snapshot)
    details = snapshot_detail_context(snapshot)
    names = [
        name
        for name, key in [("purchase", "PURCHASE"), ("sales", "SALES")]
        if snapshot.scope in ("BOTH", key)
    ]
    wb = Workbook()
    ws = wb.active
    ws.title = "采销汇总"
    ws.append([snapshot.company.name + " · 采销汇总日报"])
    ws.append(["报表日期", snapshot.report_date.isoformat()])
    business_label = " / ".join(
        "采购入库" if name == "purchase" else "销售出库" for name in names
    )
    ws.append(["来源", "金蝶 K3Cloud · 有效" + business_label + "单，退货保留负数"])
    ws.append(
        ["统计期间", data.get("period_start", "") + "—" + data.get("period_end", "")]
    )
    ws.append(["指标", "金额 / 占比"])
    for name in names:
        label = "采购" if name == "purchase" else "销售"
        ws.append(["当月" + label + "日报合计", _number(data.get(name + "_total"))])
        ws.append(["当日" + label + "金额", _number(data.get(name + "_daily"))])
    if len(names) == 2:
        ws.append(
            [
                "合计占比",
                (
                    float(data["share_total"]) / 100
                    if data.get("share_total") is not None
                    else None
                ),
            ]
        )
        ws.cell(ws.max_row, 2).number_format = "0.00%"
    if len(names) == 2:
        ws.append(
            ["口径", "合计占比 = 当月采购 ÷ 当月销售；销售为 0 时留空。不代表利润率。"]
        )
    if data.get("source") != "KINGDEE":
        ws.append(["提示", "历史快照无金蝶逐日汇总，请重新生成，不能用手工日报替代。"])
    _style(ws, 5)
    for row in ws:
        if row[0].value == "合计占比":
            row[1].number_format = "0.00%"
    for ledger in details["manual_ledgers"]:
        ws.append(["当日" + ledger["label"] + "日报提交笔数", ledger["count"]])
        ws.append(["当日" + ledger["label"] + "日报提交金额", _number(ledger["total"])])
    ws.append(
        [
            "日报提交口径",
            "仅人员填报 / 历史导入的报表当日记录；没有提交时保留表头和空状态，不以金蝶数据代填。",
        ]
    )
    for row in ws.iter_rows(min_row=6):
        if str(row[0].value).endswith("提交笔数"):
            row[1].number_format = "0"
        elif str(row[0].value).endswith("提交金额"):
            row[1].number_format = "#,##0.00;[Red]-#,##0.00"
    _create_detail_sheets(wb, snapshot, details)
    daily = _daily_sheet(wb, "当月金蝶采销", data.get("month_rows", []), names)
    _daily_sheet(wb, "近15天金蝶趋势", data.get("trend_rows", []), names)
    if daily.max_row > 1:
        chart = LineChart()
        chart.title = "当月采销对比折线图"
        chart.y_axis.title = "人民币元"
        chart.x_axis.delete = True
        chart.add_data(
            Reference(
                daily,
                min_col=2,
                max_col=1 + len(names),
                min_row=1,
                max_row=daily.max_row,
            ),
            titles_from_data=True,
        )
        chart.set_categories(
            Reference(daily, min_col=1, min_row=2, max_row=daily.max_row)
        )
        chart.height = 10
        chart.width = 22
        chart.style = 13
        for index, series in enumerate(chart.series):
            series.graphicalProperties.line.solidFill = (
                "B45309" if names[index] == "purchase" else "2563EB"
            )
        ws.add_chart(chart, "D2")
    comparison = (snapshot.data or {}).get("monthly_comparison")
    if comparison:
        title = (
            COMPARISON_SHEET
            if len(names) == 2
            else ("采购入库金额对比表" if names == ["purchase"] else "销售金额对比表")
        )
        manual = wb.create_sheet(title)
        manual.append([title])
        manual.append(["月份", f"{comparison['year']}-{comparison['month']:02d}"])
        manual.append(
            [
                "来源",
                "原有对比表 · 手工录入 / 历史导入 · 未税人民币金额；不是金蝶单据金额",
            ]
        )
        headers = ["日期"] + [
            ("采购入库金额" if name == "purchase" else "销售金额") for name in names
        ]
        if len(names) == 2:
            headers += ["每天占比"]
        headers += ["备注"]
        manual.append(headers)
        for row in comparison["rows"]:
            values = [row["date"]] + [
                _number(row.get(name + "_amount")) for name in names
            ]
            if len(names) == 2:
                values += [
                    float(row["share"]) / 100 if row.get("share") is not None else None
                ]
            manual.append(values + [""])
        totals = ["月度合计"] + [
            _number(comparison.get(name + "_total")) for name in names
        ]
        if len(names) == 2:
            totals += [
                (
                    float(comparison["share_total"]) / 100
                    if comparison.get("share_total") is not None
                    else None
                )
            ]
        manual.append(totals + [""])
        _style(manual, 4)
        if len(names) == 2:
            for row in manual.iter_rows(min_row=5):
                row[3].number_format = "0.00%"
        manual.auto_filter.ref = (
            f"A4:{manual.cell(manual.max_row-1,manual.max_column).coordinate}"
        )
    output = BytesIO()
    wb.save(output)
    return output.getvalue()

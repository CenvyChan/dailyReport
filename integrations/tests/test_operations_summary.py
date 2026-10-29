"""金蝶采销汇总、不可变月度对比快照及统一导出。"""

import json
from datetime import date
from decimal import Decimal
from io import BytesIO
from django.test import TestCase
from openpyxl import load_workbook
from core.models import Company, Customer, Supplier
from core.testing import company_a, company_b
from django.contrib.auth.models import User
from integrations.models import StockTransaction, SyncRun
from integrations.services.report_delivery import (
    generate_report_snapshot,
    render_snapshot_book_html,
    render_snapshot_html,
)
from integrations.services.operations_a2ui import (
    build_operations_a2ui_messages,
    filter_snapshot_for_contents,
)
from integrations.services.snapshot_export import export_snapshot_to_excel
from reports.tests.test_comparison import add_sale, add_purchase


class OperationsSummaryTests(TestCase):
    def setUp(self):
        self.company = company_a()
        self.other = company_b()
        self.user = User.objects.create_user("summary-user")
        self.customer = Customer.objects.create(company=self.company, name="客户")
        self.supplier = Supplier.objects.create(company=self.company, name="供应商")
        SyncRun.objects.create(
            company=self.company,
            integration="KINGDEE",
            status="SUCCESS",
            window_start=date(2026, 1, 1),
            window_end=date(2026, 10, 28),
        )
        self.stock("SALES_OUT", "100", date(2026, 10, 1), "1")
        self.stock("SALES_OUT", "200", date(2026, 10, 27), "2")
        self.stock("PURCHASE_IN", "120", date(2026, 10, 27), "3")
        self.stock("PURCHASE_IN", "-20", date(2026, 10, 26), "4")
        self.stock("SALES_OUT", "9000", date(2026, 10, 27), "5", active=False)
        self.stock("SALES_OUT", "8000", date(2026, 10, 27), "6", company=self.other)
        self.stock("SALES_OUT", "7000", date(2026, 10, 28), "7")
        add_sale(self.company, self.customer, self.user, date(2026, 10, 1), "999")
        add_purchase(self.company, self.supplier, self.user, date(2026, 10, 1), "333")

    def stock(self, kind, amount, day, fid, active=True, company=None):
        return StockTransaction.objects.create(
            company=company or self.company,
            business_type=kind,
            form_id="SAL_OUTSTOCK" if kind == "SALES_OUT" else "STK_InStock",
            fid=fid,
            bill_no=fid,
            biz_date=day,
            total_amount=amount,
            is_active=active,
        )

    def snapshot(self):
        return generate_report_snapshot(self.company, date(2026, 10, 27))

    def test_k3_summary_and_signed_two_series_do_not_use_manual_or_inactive_data(self):
        snap = self.snapshot()
        data = snap.data["operations"]
        self.assertEqual(data["source"], "KINGDEE")
        self.assertEqual(Decimal(data["purchase_total"]), Decimal("100"))
        self.assertEqual(Decimal(data["sales_total"]), Decimal("300"))
        self.assertAlmostEqual(float(data["share_total"]), 100 / 3, places=3)
        self.assertEqual(len(data["month_rows"]), 27)
        self.assertEqual(len(data["trend_rows"]), 15)
        self.assertEqual(
            Decimal(data["trend_rows"][-2]["purchase_amount"]), Decimal("-20")
        )
        self.assertEqual(
            Decimal(data["trend_rows"][-1]["sales_amount"]), Decimal("200")
        )
        model = build_operations_a2ui_messages(snap)[1]["updateDataModel"]["value"]
        self.assertEqual(model["summary"]["purchaseTotalText"], "¥100.00")
        self.assertEqual(model["summary"]["salesTotalText"], "¥300.00")
        self.assertEqual(
            {p["type"] for p in model["charts"]["month"]["data"]}, {"采购", "销售"}
        )
        self.assertTrue(any(p["y"] < 0 for p in model["charts"]["trend"]["data"]))
        for text in [
            "月度对账差额",
            "年度对账差额",
            "关联覆盖率",
            "monthDiffText",
            "yearDiffText",
        ]:
            self.assertNotIn(
                text,
                json.dumps(build_operations_a2ui_messages(snap), ensure_ascii=False),
            )

    def test_monthly_comparison_sheet_reuses_existing_report_and_does_not_query_on_export(
        self,
    ):
        snap = self.snapshot()
        add_sale(self.company, self.customer, self.user, date(2026, 10, 2), "111")
        with self.assertNumQueries(0):
            binary = export_snapshot_to_excel(snap)
        wb = load_workbook(BytesIO(binary), data_only=True)
        self.assertIn("采购入库金额与销售金额对比表", wb.sheetnames)
        self.assertNotIn("销售对账", wb.sheetnames)
        self.assertNotIn("采购对账", wb.sheetnames)
        ws = wb["采购入库金额与销售金额对比表"]
        rows = list(ws.values)
        self.assertTrue(
            any(r[0] == "2026-10-01" and r[1] == 333 and r[2] == 999 for r in rows)
        )
        self.assertTrue(
            any(r[0] == "月度合计" and r[1] == 333 and r[2] == 999 for r in rows)
        )
        self.assertTrue(wb["采销汇总"]._charts)

    def test_html_pdf_use_one_combined_report_without_reconciliation_metrics(self):
        snap = self.snapshot()
        for html in [render_snapshot_html(snap), render_snapshot_book_html(snap)]:
            self.assertIn("当月采购日报合计", html)
            self.assertIn("当月销售日报合计", html)
            self.assertIn("当月采销对比", html)
            self.assertIn("近 15 天", html)
            self.assertIn("金蝶", html)
            self.assertIn("polyline", html)
            for text in [
                "月度对账差额",
                "年度对账差额",
                "覆盖率",
                "差异率",
                "采购日报</h2>",
                "销售日报</h2>",
            ]:
                self.assertNotIn(text, html)

    def test_content_selection_never_leaks_other_business_in_charts_or_excel(self):
        selected = filter_snapshot_for_contents(self.snapshot(), ["PURCHASE"])
        model = build_operations_a2ui_messages(selected)[1]["updateDataModel"]["value"]
        self.assertEqual(
            {p["type"] for p in model["charts"]["month"]["data"]}, {"采购"}
        )
        self.assertNotIn("salesTotalText", model["summary"])
        wb = load_workbook(BytesIO(export_snapshot_to_excel(selected)))
        for ws in wb:
            self.assertFalse(
                any("销售" in str(v) for row in ws.values for v in row if v is not None)
            )

    def test_zero_sales_has_no_ratio_and_month_boundaries(self):
        snap = generate_report_snapshot(self.other, date(2026, 2, 1))
        self.assertIsNone(snap.data["operations"]["share_total"])
        self.assertEqual(len(snap.data["operations"]["month_rows"]), 1)
        self.assertEqual(len(snap.data["monthly_comparison"]["rows"]), 28)
        model = build_operations_a2ui_messages(snap)[1]["updateDataModel"]["value"]
        self.assertEqual(model["summary"]["shareText"], "—")

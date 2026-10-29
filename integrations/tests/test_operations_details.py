"""明细导出回归：无提交也保留表头，金蝶实际单据不替代手填日报。"""

from datetime import date
from decimal import Decimal
from io import BytesIO
from django.test import TestCase
from django.contrib.auth.models import User
from openpyxl import load_workbook
from core.testing import company_a, company_b
from core.models import Customer, Supplier
from reports.tests.test_comparison import add_sale, add_purchase
from integrations.models import StockTransaction, StockTransactionLine, SyncRun
from integrations.services.report_delivery import (
    generate_report_snapshot,
    render_snapshot_book_html,
    render_snapshot_html,
)
from integrations.services.snapshot_export import export_snapshot_to_excel
from integrations.services.operations_a2ui import filter_snapshot_for_contents


class OperationsDetailsTests(TestCase):
    def setUp(self):
        self.company = company_a()
        self.other = company_b()
        self.user = User.objects.create_user("detail-user", first_name="提报员")
        self.customer = Customer.objects.create(company=self.company, name="手填客户")
        self.supplier = Supplier.objects.create(company=self.company, name="手填供应商")
        SyncRun.objects.create(
            company=self.company,
            integration="KINGDEE",
            status="SUCCESS",
            window_start=date(2026, 1, 1),
            window_end=date(2026, 10, 28),
        )
        self.stock("PURCHASE_IN", "120", "RK001")
        self.stock("SALES_OUT", "200", "CK001")

    def stock(self, kind, amount, bill, **overrides):
        values = dict(
            company=self.company,
            business_type=kind,
            biz_date=date(2026, 10, 27),
            form_id="SAL_OUTSTOCK" if kind == "SALES_OUT" else "STK_InStock",
            fid=bill,
            bill_no=bill,
            k3_party_name="金蝶客户" if kind == "SALES_OUT" else "金蝶供应商",
            total_amount=amount,
            is_active=True,
        )
        values.update(overrides)
        return StockTransaction.objects.create(**values)

    def snapshot(self):
        return generate_report_snapshot(
            self.company, date(2026, 10, 27), force_regenerate=True
        )

    def test_empty_manual_export_preserves_columns_and_explicit_empty_state(self):
        snap = self.snapshot()
        wb = load_workbook(BytesIO(export_snapshot_to_excel(snap)), data_only=True)
        sheet = wb["日报提交明细"]
        self.assertEqual(
            list(sheet.values)[3],
            (
                "业务类型",
                "日期",
                "客户/供应商",
                "提报人",
                "数量",
                "单位",
                "日报金额（元）",
            ),
        )
        self.assertIn("暂无提交记录", sheet["A5"].value)
        self.assertIsNone(sheet["E5"].value)
        self.assertEqual(sheet["G6"].value, 0)
        html = render_snapshot_book_html(snap)
        self.assertIn("cover-frame", html)
        for title in [
            "经营总览与关键指标",
            "销售日报明细账簿",
            "采购日报明细账簿",
            "客户名称",
            "供应商名称",
            "提报人",
            "暂无提交记录",
        ]:
            self.assertIn(title, html)
        self.assertIn("当日销售日报提交笔数", str(list(wb["采销汇总"].values)))
        self.assertNotIn("月度对账差额", html)
        self.assertNotIn("年度对账差额", html)

    def test_submitted_manual_rows_are_frozen_and_export_does_not_query(self):
        add_sale(self.company, self.customer, self.user, date(2026, 10, 27), "12.34")
        add_purchase(self.company, self.supplier, self.user, date(2026, 10, 27), "5.67")
        snap = self.snapshot()
        add_sale(self.company, self.customer, self.user, date(2026, 10, 27), "99")
        with self.assertNumQueries(0):
            wb = load_workbook(BytesIO(export_snapshot_to_excel(snap)), data_only=True)
            html = render_snapshot_book_html(snap)
        rows = list(wb["日报提交明细"].values)
        records = [r for r in rows[4:] if r[0] in ("采购", "销售")]
        self.assertEqual(len(records), 2)
        self.assertEqual({r[2] for r in records}, {"手填客户", "手填供应商"})
        self.assertEqual({r[3] for r in records}, {"提报员"})
        self.assertEqual(sum(Decimal(str(r[-1])) for r in records), Decimal("18.01"))
        self.assertEqual(rows[-1][-1], 18.01)
        self.assertIn("12.34", html)
        self.assertIn("5.67", html)

    def test_actual_kingdee_documents_match_day_company_validity_and_signed_amounts(
        self,
    ):
        self.stock("PURCHASE_IN", "-20", "RK-RETURN")
        self.stock("SALES_OUT", "999", "OTHER", company=self.other)
        self.stock("SALES_OUT", "888", "VOID", is_active=False)
        self.stock("SALES_OUT", "777", "PREVIOUS", biz_date=date(2026, 10, 26))
        snap = self.snapshot()
        docs = snap.data["operations"]["daily_documents"]
        self.assertEqual({r["bill_no"] for r in docs}, {"RK001", "CK001", "RK-RETURN"})
        self.assertEqual(
            sum(Decimal(r["amount"]) for r in docs if r["business_type"] == "PURCHASE"),
            Decimal("100"),
        )
        wb = load_workbook(BytesIO(export_snapshot_to_excel(snap)), data_only=True)
        records = list(wb["当日金蝶单据明细"].values)[4:-1]
        self.assertEqual(len(records), 3)
        self.assertIn(-20, [r[-1] for r in records])
        html = render_snapshot_html(snap)
        self.assertIn("当日金蝶实际单据明细", html)
        for bill in ["RK001", "CK001", "RK-RETURN"]:
            self.assertIn(bill, html)
        self.assertNotIn("PREVIOUS", html)

    def test_single_business_projection_never_leaks_manual_or_kingdee_rows(self):
        add_sale(self.company, self.customer, self.user, date(2026, 10, 27), "12.34")
        add_purchase(self.company, self.supplier, self.user, date(2026, 10, 27), "5.67")
        snap = filter_snapshot_for_contents(self.snapshot(), ["PURCHASE"])
        self.assertTrue(
            all(
                r["business_type"] == "PURCHASE"
                for r in snap.data["operations"]["daily_documents"]
            )
        )
        html = render_snapshot_book_html(snap)
        for word in ["CK001", "手填客户", "销售日报明细账簿"]:
            self.assertNotIn(word, html)
        wb = load_workbook(BytesIO(export_snapshot_to_excel(snap)), data_only=True)
        text = str([list(ws.values) for ws in wb])
        self.assertNotIn("CK001", text)
        self.assertNotIn("手填客户", text)

    def test_all_kingdee_documents_exported_and_paginated_without_first_100_cutoff(
        self,
    ):
        for n in range(105):
            self.stock("SALES_OUT", "1", "CK%04d" % (n + 100))
        snap = self.snapshot()
        self.assertEqual(len(snap.data["operations"]["daily_documents"]), 107)
        html = render_snapshot_book_html(snap)
        self.assertIn("CK0204", html)
        self.assertIn("明细续页", html)
        wb = load_workbook(BytesIO(export_snapshot_to_excel(snap)), data_only=True)
        self.assertEqual(len(list(wb["当日金蝶单据明细"].values)[4:-1]), 107)

    def test_quantities_and_mixed_units_are_frozen_and_summed_without_conversion(self):
        purchase = StockTransaction.objects.get(bill_no="RK001")
        purchase.total_quantity = Decimal("12.345678")
        purchase.save(update_fields=["total_quantity"])
        sale = StockTransaction.objects.get(bill_no="CK001")
        sale.total_quantity = Decimal("8")
        sale.save(update_fields=["total_quantity"])
        StockTransactionLine.objects.create(
            transaction=purchase,
            entry_id="1",
            quantity="12.345678",
            unit="千克",
            amount="120",
        )
        StockTransactionLine.objects.create(
            transaction=sale, entry_id="2", quantity="8", unit="件", amount="200"
        )
        returned = self.stock(
            "PURCHASE_IN", "-20", "RK-RETURN", total_quantity="-2.345678"
        )
        StockTransactionLine.objects.create(
            transaction=returned,
            entry_id="3",
            quantity="-2.345678",
            unit="千克",
            amount="-20",
        )
        snap = self.snapshot()
        purchase.total_quantity = Decimal("999")
        purchase.save(update_fields=["total_quantity"])
        with self.assertNumQueries(0):
            wb = load_workbook(BytesIO(export_snapshot_to_excel(snap)), data_only=True)
            html = render_snapshot_book_html(snap)
        sheet = wb["当日金蝶单据明细"]
        self.assertEqual(
            list(sheet.values)[3],
            (
                "业务类型",
                "业务日期",
                "单据编号",
                "客户/供应商",
                "数量",
                "单位",
                "金蝶金额（元）",
            ),
        )
        rows = list(sheet.values)
        self.assertEqual(rows[-1][4], 18)
        self.assertEqual(rows[-1][-1], 300)
        self.assertIn("混合单位", rows[-1][5])
        self.assertIn("12.345678", html)
        self.assertIn("-2.345678", html)
        self.assertIn("千克", html)
        self.assertIn("数量合计", html)
        self.assertIn("金额合计", html)
        self.assertIn("单据编号", html)

    def test_manual_quantity_and_empty_quantity_totals(self):
        sale = add_sale(
            self.company, self.customer, self.user, date(2026, 10, 27), "12.34"
        )
        sale.quantity = Decimal("1.250")
        sale.save(update_fields=["quantity"])
        purchase = add_purchase(
            self.company, self.supplier, self.user, date(2026, 10, 27), "5.67"
        )
        purchase.quantity = Decimal("2.500")
        purchase.save(update_fields=["quantity"])
        snap = self.snapshot()
        wb = load_workbook(BytesIO(export_snapshot_to_excel(snap)), data_only=True)
        self.assertEqual(list(wb["日报提交明细"].values)[-1][4], 3.75)
        self.assertEqual(list(wb["日报提交明细"].values)[-1][-1], 18.01)
        html = render_snapshot_book_html(snap)
        self.assertIn("1.25", html)
        self.assertIn("2.5", html)

    def test_legacy_snapshot_without_frozen_documents_is_explicit_not_faked(self):
        snap = self.snapshot()
        snap.data["operations"].pop("daily_documents")
        html = render_snapshot_book_html(snap)
        self.assertIn("历史快照未冻结金蝶单据明细，请重新生成", html)
        wb = load_workbook(BytesIO(export_snapshot_to_excel(snap)), data_only=True)
        self.assertIn("历史快照未冻结", wb["当日金蝶单据明细"]["A5"].value)

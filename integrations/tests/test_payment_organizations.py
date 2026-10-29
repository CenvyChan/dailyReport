from datetime import date
from decimal import Decimal
from io import StringIO
from unittest.mock import patch

from django.core.management import call_command
from django.test import TestCase

from core.models import Company
from integrations.models import DingtalkApp, DingtalkApprovalInstance, DingtalkProcessConfig
from integrations.services.payment_report import _extract_form_value


class PaymentFormCompatibilityTests(TestCase):
    def test_dictionary_form_preserves_recipient(self):
        self.assertEqual(_extract_form_value({"收款人": "供应商甲"}, "收款人"), "供应商甲")

    def test_list_form_preserves_recipient(self):
        self.assertEqual(_extract_form_value([{"name": "收款人", "value": "供应商甲"}], "收款人"), "供应商甲")


from datetime import datetime, time
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import Mock
from django.core.management.base import CommandError
from django.utils import timezone

from integrations.models import DingtalkCompanyMapping, DingtalkGroupConfig, PaymentReportDispatch
from integrations.services.payment_organizations import (PAYMENT_REPORT_COMPANY_APPLICATION_NAMES, save_approval, flatten_form)
from integrations.services.payment_report import analyze_payment_applications, get_previous_day_range, get_last_week_range


class PaymentOrganizationTests(TestCase):
    def setUp(self):
        self.source = Company.objects.create(code="SOURCE", name="来源公司")
        self.a = Company.objects.create(code="FNS", name="飞诺斯电子科技有限公司")
        self.b = Company.objects.create(code="NBHH", name="宁波汇慧有限公司")
        self.app = DingtalkApp.objects.create(company=self.source, app_key="k", app_secret="s", robot_code="r")
        self.config = DingtalkProcessConfig.objects.create(app=self.app, process_code="PROC-PAY", name="付款申请", amount_component="付款金额")
        for company in (self.a, self.b):
            DingtalkCompanyMapping.objects.create(
                app=self.app,
                application_company=PAYMENT_REPORT_COMPANY_APPLICATION_NAMES[company.code],
                company=company,
            )

    def raw(self, company=None, business="BIZ-1"):
        company = company or self.a
        application_company = PAYMENT_REPORT_COMPANY_APPLICATION_NAMES.get(company.code, company.name)
        return {"business_id": business, "status": "COMPLETED", "originator_name": "张三",
                "create_time": "2026-10-26 10:00:00", "form_component_values": [
                    {"name": "申请公司", "value": application_company},
                    {"name": "付款金额", "value": "123.45"},
                    {"name": "收款人", "value": "供应商甲"}]}

    def test_shared_application_routes_to_seven_separate_companies(self):
        companies = [self.a, self.b]
        for number in range(5):
            company = Company.objects.create(code=f"PAY-{number}", name=f"其他申请公司{number}")
            DingtalkCompanyMapping.objects.create(app=self.app, application_company=company.name, company=company)
            companies.append(company)
        for number, company in enumerate(companies):
            for repeat in range(2):
                save_approval(self.app, self.config, self.raw(company, f"BIZ-{number}"), f"PID-{number}")
        self.assertEqual(DingtalkApprovalInstance.objects.count(), 7)
        period = get_previous_day_range(date(2026, 10, 27))
        for company in companies:
            data = analyze_payment_applications(company, period["start"], period["end"], self.config.process_code)
            self.assertEqual(data["summary"]["total_count"], 1)
            self.assertEqual(data["summary"]["total_amount"], Decimal("123.45"))
            self.assertEqual(data["by_recipient"][0]["recipient"], "供应商甲")
        self.assertFalse(DingtalkApprovalInstance.objects.filter(company=self.source).exists())

    def test_same_amount_and_recipient_are_not_deduplicated(self):
        for pid in ("PID-1", "PID-2"):
            save_approval(self.app, self.config, self.raw(business=pid), pid)
        self.assertEqual(DingtalkApprovalInstance.objects.count(), 2)

    def test_period_count_includes_all_current_statuses_and_rejected_results(self):
        period = get_previous_day_range(date(2026, 10, 27))
        for status in ("NEW", "RUNNING", "COMPLETED", "TERMINATED"):
            raw = self.raw(business=f"BIZ-{status}")
            raw["status"] = status
            save_approval(self.app, self.config, raw, f"PID-{status}")
        rejected = self.raw(business="BIZ-REJECTED")
        rejected["result"] = "refuse"
        save_approval(self.app, self.config, rejected, "PID-REJECTED")
        for number, created_at in enumerate(("2026-10-25 23:59:59", "2026-10-27 00:00:00")):
            raw = self.raw(business=f"BIZ-OUTSIDE-{number}")
            raw["create_time"] = created_at
            save_approval(self.app, self.config, raw, f"PID-OUTSIDE-{number}")

        data = analyze_payment_applications(
            self.a, period["start"], period["end"], self.config.process_code
        )

        self.assertEqual(data["summary"]["total_count"], 5)
        self.assertEqual(data["summary"]["total_amount"], Decimal("617.25"))
        self.assertEqual(data["summary"]["by_status"], {
            "NEW": 1, "RUNNING": 1, "COMPLETED": 2, "TERMINATED": 1,
        })

    def test_unknown_company_fails_instead_of_using_source_company(self):
        raw = self.raw()
        raw["form_component_values"][0]["value"] = "未配置的第八家公司"
        with self.assertRaisesMessage(ValueError, "尚未映射"):
            save_approval(self.app, self.config, raw, "PID-1")
        self.assertFalse(DingtalkApprovalInstance.objects.exists())

    def test_company_change_moves_existing_instance_without_creating_copy(self):
        original, _, _ = save_approval(self.app, self.config, self.raw(self.a), "PID-1")
        updated, created, _ = save_approval(self.app, self.config, self.raw(self.b), "PID-1")
        self.assertFalse(created)
        self.assertEqual(updated.pk, original.pk)
        self.assertEqual(updated.company, self.b)
        self.assertEqual(DingtalkApprovalInstance.objects.count(), 1)

    def test_proven_business_id_alias_requires_repair_and_then_merges(self):
        raw = self.raw()
        save_approval(self.app, self.config, raw, "PID-1")
        DingtalkApprovalInstance.objects.create(company=self.source, process_code=self.config.process_code,
            process_instance_id="BIZ-1", raw=raw, form_values=flatten_form(raw["form_component_values"]), amount=Decimal("123.45"))
        with self.assertRaisesMessage(ValueError, "重复记录"):
            save_approval(self.app, self.config, raw, "PID-1")
        save_approval(self.app, self.config, raw, "PID-1", repair=True)
        self.assertEqual(DingtalkApprovalInstance.objects.count(), 1)
        self.assertEqual(DingtalkApprovalInstance.objects.get().company, self.a)

    def test_all_three_sync_entrypoints_use_same_identity(self):
        from integrations.services.dingtalk_sync import _upsert_instance
        from integrations.management.commands.sync_payment_approvals import Command as PaymentSync
        from integrations.management.commands.sync_running_approvals import Command as RunningSync
        _upsert_instance(self.source, self.config, self.raw(), None, "PID-1")
        client = Mock()
        client.get_instance.return_value = self.raw()
        PaymentSync()._sync_instances(self.source, self.config, client, ["PID-1"], False, True)
        RunningSync()._save_instance(self.source, self.config, "PID-1", self.raw())
        self.assertEqual(DingtalkApprovalInstance.objects.count(), 1)
        self.assertEqual(DingtalkApprovalInstance.objects.get().company, self.a)

    def test_reconcile_keeps_latest_synced_payload_for_explicit_duplicates(self):
        from datetime import timedelta

        for newest_id in ("PID-RECENT", "BIZ-RECENT"):
            with self.subTest(newest_id=newest_id), TemporaryDirectory() as directory:
                DingtalkApprovalInstance.objects.all().delete()
                for instance_id in ("PID-RECENT", "BIZ-RECENT"):
                    raw = self.raw(business="BIZ-RECENT")
                    raw["process_instance_id"] = "PID-RECENT"
                    raw["status"] = "COMPLETED" if instance_id == newest_id else "RUNNING"
                    raw["result"] = "agree" if instance_id == newest_id else ""
                    row = DingtalkApprovalInstance.objects.create(
                        company=self.source, process_code=self.config.process_code,
                        process_instance_id=instance_id, raw=raw,
                        form_values=flatten_form(raw["form_component_values"]),
                        amount=Decimal("123.45"),
                    )
                    updated = timezone.now() - timedelta(days=0 if instance_id == newest_id else 2)
                    DingtalkApprovalInstance.objects.filter(pk=row.pk).update(updated_at=updated)
                call_command("reconcile_payment_organizations", app_id=self.app.pk,
                             apply=True, backup_dir=directory, stdout=StringIO())
                self.assertEqual(DingtalkApprovalInstance.objects.count(), 1)
                row = DingtalkApprovalInstance.objects.get()
                self.assertEqual(row.process_instance_id, "PID-RECENT")
                self.assertEqual(row.status, "COMPLETED")
                self.assertEqual(row.result, "agree")

    def test_reconcile_maps_full_form_names_to_existing_project_companies(self):
        for company in (self.a, self.b):
            form_name = PAYMENT_REPORT_COMPANY_APPLICATION_NAMES[company.code]
            DingtalkCompanyMapping.objects.filter(
                app=self.app,
                application_company=form_name,
            ).update(company=self.source)
            instance_id = f"PID-{company.code}"
            save_approval(self.app, self.config, self.raw(company, instance_id), instance_id)

        with TemporaryDirectory() as directory:
            call_command(
                "reconcile_payment_organizations",
                app_id=self.app.pk,
                apply=True,
                backup_dir=directory,
                stdout=StringIO(),
            )

        mappings = {
            mapping.application_company: mapping.company
            for mapping in self.app.company_mappings.select_related("company")
        }
        self.assertEqual(mappings[PAYMENT_REPORT_COMPANY_APPLICATION_NAMES["FNS"]], self.a)
        self.assertEqual(mappings[PAYMENT_REPORT_COMPANY_APPLICATION_NAMES["NBHH"]], self.b)
        self.assertFalse(Company.objects.filter(name=PAYMENT_REPORT_COMPANY_APPLICATION_NAMES["FNS"]).exists())
        self.assertFalse(Company.objects.filter(name=PAYMENT_REPORT_COMPANY_APPLICATION_NAMES["NBHH"]).exists())
        self.assertEqual(set(DingtalkApprovalInstance.objects.values_list("company_id", flat=True)), {self.a.pk, self.b.pk})

    def test_reconcile_does_not_create_target_company_when_code_missing(self):
        form_name = PAYMENT_REPORT_COMPANY_APPLICATION_NAMES["FNS"]
        DingtalkCompanyMapping.objects.filter(
            app=self.app, application_company=form_name
        ).update(company=self.source)
        save_approval(self.app, self.config, self.raw(self.a, "PID-FNS"), "PID-FNS")
        DingtalkCompanyMapping.objects.filter(app=self.app, application_company=form_name).delete()
        self.a.delete()

        with TemporaryDirectory() as directory:
            with self.assertRaisesMessage(CommandError, "项目组织代码 FNS 不存在或未启用"):
                call_command(
                    "reconcile_payment_organizations",
                    app_id=self.app.pk,
                    apply=True,
                    create_missing=True,
                    backup_dir=directory,
                    stdout=StringIO(),
                )

        self.assertFalse(Company.objects.filter(name=form_name).exists())

    def test_audit_is_read_only_and_apply_backs_up_before_reclassifying(self):
        raw = self.raw()
        DingtalkApprovalInstance.objects.create(company=self.source, process_code=self.config.process_code,
            process_instance_id="PID-1", raw=raw, form_values=raw["form_component_values"], amount=Decimal("123.45"))
        out = StringIO()
        with TemporaryDirectory() as directory:
            call_command("reconcile_payment_organizations", app_id=self.app.pk, backup_dir=directory, stdout=out)
            self.assertEqual(DingtalkApprovalInstance.objects.get().company, self.source)
            self.assertEqual(list(Path(directory).iterdir()), [])
            call_command("reconcile_payment_organizations", app_id=self.app.pk, apply=True, backup_dir=directory, stdout=out)
            self.assertEqual(DingtalkApprovalInstance.objects.get().company, self.a)
            self.assertEqual(len(list(Path(directory).glob("*.json"))), 1)
            call_command("loaddata", str(next(Path(directory).glob("*.json"))), stdout=StringIO())
            self.assertEqual(DingtalkApprovalInstance.objects.get().company, self.source)

    def test_local_repair_does_not_reinsert_business_alias(self):
        raw = self.raw()
        for pid in ("PID-1", "BIZ-1"):
            DingtalkApprovalInstance.objects.create(company=self.source, process_code=self.config.process_code,
                process_instance_id=pid, raw=raw, form_values=raw["form_component_values"], amount=Decimal("123.45"))
        with TemporaryDirectory() as directory:
            call_command("reconcile_payment_organizations", app_id=self.app.pk, apply=True, backup_dir=directory, stdout=StringIO())
        self.assertEqual(list(DingtalkApprovalInstance.objects.values_list("process_instance_id", flat=True)), ["PID-1"])

    def test_refresh_uses_server_instance_ids_then_removes_proven_alias(self):
        raw = self.raw()
        DingtalkApprovalInstance.objects.create(company=self.source, process_code=self.config.process_code,
            process_instance_id="BIZ-1", raw=raw, form_values=flatten_form(raw["form_component_values"]), amount=Decimal("123.45"))
        client = Mock()
        client.list_instance_ids.return_value = ["PID-1"]
        client.get_instance.return_value = raw
        with TemporaryDirectory() as directory, patch("integrations.management.commands.reconcile_payment_organizations.DingtalkClient", return_value=client):
            call_command("reconcile_payment_organizations", app_id=self.app.pk, apply=True, refresh=True,
                         start="2026-10-01", end="2026-10-27", backup_dir=directory, stdout=StringIO())
        self.assertEqual(list(DingtalkApprovalInstance.objects.values_list("process_instance_id", flat=True)), ["PID-1"])

    def test_company_groups_share_source_app_and_dispatch_separately_once(self):
        groups = []
        for company in (self.a, self.b):
            groups.append(DingtalkGroupConfig.objects.create(company=company, app=self.app, name=company.name,
                open_conversation_id=f"cid-{company.code}", send_at=time(9), content_types=[], payment_report_types=["day", "week"]))
        other = Company.objects.create(code="PAY-OTHER", name="其他付款公司")
        DingtalkCompanyMapping.objects.create(app=self.app, application_company=other.name, company=other)
        groups.append(DingtalkGroupConfig.objects.create(company=other, app=self.app, name=other.name,
            open_conversation_id="cid-other", send_at=time(9), content_types=[], payment_report_types=["day", "week"]))
        monday = timezone.make_aware(datetime(2026, 10, 26, 9, 10))
        with patch("integrations.management.commands.send_scheduled_payment_reports.timezone.localtime", return_value=monday), patch("integrations.management.commands.send_scheduled_payment_reports.call_command") as send:
            for _ in range(2):
                call_command("send_scheduled_payment_reports", stdout=StringIO())
        self.assertEqual(send.call_count, 4)
        self.assertEqual({call.kwargs["company"] for call in send.call_args_list}, {self.a.code, self.b.code})
        self.assertEqual({call.kwargs["group_id"] for call in send.call_args_list},
            {g.open_conversation_id for g in groups if g.company.code in ("FNS", "NBHH")})
        self.assertEqual(PaymentReportDispatch.objects.filter(status="SENT").count(), 4)

    def test_one_group_can_push_selected_companies_to_same_chat(self):
        group = DingtalkGroupConfig.objects.create(
            company=None,
            app=self.app,
            name="共享付款群",
            open_conversation_id="cid-shared",
            send_at=time(9),
            content_types=[],
            payment_report_types=["day", "week"],
        )
        group.companies.set([self.a, self.b])
        monday = timezone.make_aware(datetime(2026, 10, 26, 9, 10))
        with patch(
            "integrations.management.commands.send_scheduled_payment_reports.timezone.localtime",
            return_value=monday,
        ), patch(
            "integrations.management.commands.send_scheduled_payment_reports.call_command"
        ) as send:
            call_command("send_scheduled_payment_reports", stdout=StringIO())

        self.assertEqual(send.call_count, 4)
        self.assertEqual({call.kwargs["company"] for call in send.call_args_list}, {self.a.code, self.b.code})
        self.assertEqual({call.kwargs["group_id"] for call in send.call_args_list}, {"cid-shared"})
        self.assertEqual(
            set(PaymentReportDispatch.objects.values_list("company__code", flat=True)),
            {self.a.code, self.b.code},
        )

    def test_group_only_pushes_companies_that_are_selected(self):
        group = DingtalkGroupConfig.objects.create(
            company=None,
            app=self.app,
            name="仅飞诺斯付款群",
            open_conversation_id="cid-fns-only",
            send_at=time(9),
            content_types=[],
            payment_report_types=["day"],
        )
        group.companies.set([self.a])
        now = timezone.make_aware(datetime(2026, 10, 27, 9, 10))
        with patch(
            "integrations.management.commands.send_scheduled_payment_reports.timezone.localtime",
            return_value=now,
        ), patch(
            "integrations.management.commands.send_scheduled_payment_reports.call_command"
        ) as send:
            call_command("send_scheduled_payment_reports", stdout=StringIO())

        self.assertEqual(send.call_count, 1)
        self.assertEqual(send.call_args.kwargs["company"], self.a.code)

    def test_shared_application_group_is_allowed_in_admin(self):
        from integrations.admin import DingtalkGroupConfigForm
        form = DingtalkGroupConfigForm(data={"companies":[self.a.pk, self.b.pk], "app":self.app.pk, "name":"付款群",
            "open_conversation_id":"cid-test", "send_at":"09:00", "content_types":[], "payment_report_types":["day", "week"], "is_active":True})
        self.assertTrue(form.is_valid(), form.errors)

    def test_dry_run_exports_one_card_and_excel_per_application_company(self):
        from openpyxl import load_workbook
        import json
        for number, company in enumerate((self.a, self.b)):
            save_approval(self.app, self.config, self.raw(company, str(number)), f"PID-{number}")
        other = Company.objects.create(code="PAY-OTHER", name="其他付款公司")
        DingtalkCompanyMapping.objects.create(app=self.app, application_company=other.name, company=other)
        save_approval(self.app, self.config, self.raw(other, "2"), "PID-OTHER")
        with TemporaryDirectory() as directory:
            call_command("send_payment_report", "day", base_date="2026-10-27", dry_run=True, output_dir=directory, stdout=StringIO())
            self.assertEqual(len(list(Path(directory).glob("*.xlsx"))), 2)
            for company in (self.a, self.b):
                cards = list(Path(directory).glob(f"*{company.code}_*.a2ui.json"))
                self.assertEqual(len(cards), 1)
                model = json.loads(cards[0].read_text(encoding="utf-8"))[1]["updateDataModel"]["value"]
                self.assertEqual(model["companyName"], PAYMENT_REPORT_COMPANY_APPLICATION_NAMES[company.code])
                self.assertEqual(model["summary"]["totalCount"], 1)
                self.assertEqual(model["summary"]["totalAmountText"], "¥123.45")
                book = load_workbook(next(Path(directory).glob(f"*{company.code}_*.xlsx")))
                rows = dict(book["汇总统计"].iter_rows(min_row=2, max_col=2, values_only=True))
                self.assertEqual(rows["申请公司"], PAYMENT_REPORT_COMPANY_APPLICATION_NAMES[company.code])
                self.assertEqual(rows["总申请数"], 1)
                book.close()


    def test_direct_report_rejects_non_target_company(self):
        other = Company.objects.create(code="PAY-OTHER", name="其他付款公司")
        DingtalkCompanyMapping.objects.create(app=self.app, application_company=other.name, company=other)
        with TemporaryDirectory() as directory:
            with self.assertRaises(CommandError):
                call_command("send_payment_report", "day", company=other.code, base_date="2026-10-27",
                    dry_run=True, output_dir=directory, stdout=StringIO())

    def test_failed_dispatch_is_not_automatically_replayed(self):
        DingtalkGroupConfig.objects.create(company=self.a, app=self.app, name="甲群",
            open_conversation_id="cid-a", send_at=time(9), content_types=[], payment_report_types=["day"])
        now = timezone.make_aware(datetime(2026, 10, 27, 9, 10))
        with patch("integrations.management.commands.send_scheduled_payment_reports.timezone.localtime", return_value=now), patch("integrations.management.commands.send_scheduled_payment_reports.call_command", side_effect=CommandError("附件发送失败")) as send:
            with self.assertRaises(CommandError):
                call_command("send_scheduled_payment_reports", stdout=StringIO(), stderr=StringIO())
            call_command("send_scheduled_payment_reports", stdout=StringIO(), stderr=StringIO())
        self.assertEqual(send.call_count, 1)
        self.assertEqual(PaymentReportDispatch.objects.get().status, "FAILED")

    def test_weekly_schedule_waits_until_monday_and_dry_run_has_no_side_effects(self):
        DingtalkGroupConfig.objects.create(company=self.a, app=self.app, name="甲群",
            open_conversation_id="cid-a", send_at=time(9), content_types=[], payment_report_types=["week"])
        now = timezone.make_aware(datetime(2026, 10, 27, 10))
        with patch("integrations.management.commands.send_scheduled_payment_reports.timezone.localtime", return_value=now), patch("integrations.management.commands.send_scheduled_payment_reports.call_command") as send:
            call_command("send_scheduled_payment_reports", stdout=StringIO())
            call_command("send_scheduled_payment_reports", dry_run=True, stdout=StringIO())
        send.assert_not_called()
        self.assertFalse(PaymentReportDispatch.objects.exists())

    def test_unmapped_company_repair_rolls_back_and_can_explicitly_create_company(self):
        raw = self.raw()
        raw["form_component_values"][0]["value"] = "新增申请公司"
        DingtalkApprovalInstance.objects.create(company=self.source, process_code=self.config.process_code,
            process_instance_id="PID-NEW", raw=raw, form_values=raw["form_component_values"], amount=Decimal("123.45"))
        with TemporaryDirectory() as directory:
            with self.assertRaises(CommandError):
                call_command("reconcile_payment_organizations", app_id=self.app.pk, apply=True,
                             backup_dir=directory, stdout=StringIO())
            self.assertFalse(Company.objects.filter(name="新增申请公司").exists())
            call_command("reconcile_payment_organizations", app_id=self.app.pk, apply=True, create_missing=True,
                         backup_dir=directory, stdout=StringIO())
        self.assertEqual(DingtalkApprovalInstance.objects.get().company.name, "新增申请公司")

    def test_missing_application_company_cannot_be_silently_assigned(self):
        raw = self.raw()
        raw["form_component_values"] = raw["form_component_values"][1:]
        with self.assertRaisesMessage(ValueError, "缺少申请公司"):
            save_approval(self.app, self.config, raw, "PID-MISSING")

    def test_configured_group_can_receive_selected_company_report(self):
        group = DingtalkGroupConfig.objects.create(
            company=None,
            app=self.app,
            name="共享群",
            open_conversation_id="cid-shared",
            send_at=time(9),
            content_types=[],
            payment_report_types=["day"],
        )
        group.companies.set([self.a, self.b])
        with TemporaryDirectory() as directory:
            call_command(
                "send_payment_report",
                "day",
                company=self.a.code,
                group_id="cid-shared",
                base_date="2026-10-27",
                dry_run=True,
                output_dir=directory,
                stdout=StringIO(),
            )

    def test_reports_reject_other_applicant_even_if_mapped_to_target_company(self):
        other_name = "其他付款公司"
        DingtalkCompanyMapping.objects.create(app=self.app, application_company=other_name, company=self.a)
        raw = self.raw(self.a, "PID-OTHER-APPLICANT")
        raw["form_component_values"][0]["value"] = other_name
        save_approval(self.app, self.config, raw, "PID-OTHER-APPLICANT")
        period = get_previous_day_range(date(2026, 10, 27))
        with self.assertRaisesMessage(ValueError, "不在付款报表范围"):
            analyze_payment_applications(self.a, period["start"], period["end"], self.config.process_code)

    def test_reports_reject_unreclassified_rows_with_another_application_company(self):
        raw = self.raw(self.b)
        DingtalkApprovalInstance.objects.create(company=self.a, process_code=self.config.process_code,
            process_instance_id="PID-BAD", form_values=raw["form_component_values"], amount=Decimal("123.45"),
            create_time=timezone.make_aware(datetime(2026, 10, 26, 10)))
        period = get_previous_day_range(date(2026, 10, 27))
        with self.assertRaisesMessage(ValueError, "不在付款报表范围"):
            analyze_payment_applications(self.a, period["start"], period["end"], self.config.process_code)


    def test_payment_only_group_does_not_receive_business_daily_snapshot(self):
        from types import SimpleNamespace
        from integrations.services.report_delivery import _send_to_dingtalk_channels
        DingtalkGroupConfig.objects.create(company=self.a, app=self.app, name="仅付款群",
            open_conversation_id="cid-a", send_at=time(9), content_types=[], payment_report_types=["day"])
        snapshot = SimpleNamespace(company=self.a, scope="BOTH")
        self.assertEqual(_send_to_dingtalk_channels(snapshot, send_now=True), [])

    def test_legacy_company_is_not_sent_when_deselected(self):
        group = DingtalkGroupConfig.objects.create(
            company=self.a, app=self.app, name="shared", open_conversation_id="cid-shared",
            send_at=time(9), content_types=[], payment_report_types=["day"],
        )
        group.companies.set([self.b])
        now = timezone.make_aware(datetime(2026, 10, 26, 9, 10))
        with patch("integrations.management.commands.send_scheduled_payment_reports.timezone.localtime", return_value=now), patch("integrations.management.commands.send_scheduled_payment_reports.call_command") as send:
            call_command("send_scheduled_payment_reports", stdout=StringIO())
        self.assertEqual([call.kwargs["company"] for call in send.call_args_list], [self.b.code])

    def test_company_option_limits_shared_group_to_one_company(self):
        group = DingtalkGroupConfig.objects.create(
            company=self.a, app=self.app, name="shared", open_conversation_id="cid-shared",
            send_at=time(9), content_types=[], payment_report_types=["day"],
        )
        group.companies.set([self.a, self.b])
        now = timezone.make_aware(datetime(2026, 10, 26, 9, 10))
        with patch("integrations.management.commands.send_scheduled_payment_reports.timezone.localtime", return_value=now), patch("integrations.management.commands.send_scheduled_payment_reports.call_command") as send:
            call_command("send_scheduled_payment_reports", company=self.b.code, stdout=StringIO())
        self.assertEqual([call.kwargs["company"] for call in send.call_args_list], [self.b.code])

    def test_report_rejects_configured_group_when_company_not_selected(self):
        group = DingtalkGroupConfig.objects.create(
            company=self.a, app=self.app, name="shared", open_conversation_id="cid-shared",
            send_at=time(9), content_types=[], payment_report_types=["week"],
        )
        group.companies.set([self.b])
        save_approval(self.app, self.config, self.raw(self.a), "PID-1")
        with TemporaryDirectory() as directory, patch("integrations.management.commands.send_payment_report.DingtalkClient") as client:
            with self.assertRaises(CommandError):
                call_command("send_payment_report", "week", company=self.a.code, app_id=self.app.pk,
                             group_id="cid-shared", base_date="2026-10-27", output_dir=directory,
                             stdout=StringIO(), stderr=StringIO())
        client.assert_not_called()

    def test_form_rejects_removing_all_companies_from_legacy_group(self):
        from integrations.admin import DingtalkGroupConfigForm
        group = DingtalkGroupConfig.objects.create(
            company=self.a, app=self.app, name="shared", open_conversation_id="cid-shared",
            send_at=time(9), content_types=[], payment_report_types=["week"],
        )
        group.companies.set([self.a])
        form = DingtalkGroupConfigForm(data={"companies": [], "app": self.app.pk, "name": "shared",
            "open_conversation_id": "cid-shared", "send_at": "09:00", "is_active": True,
            "content_types": [], "payment_report_types": ["week"]}, instance=group)
        self.assertFalse(form.is_valid())
